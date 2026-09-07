# -*- coding: utf-8 -*-
"""简历文件读取：从 .docx / .pdf 提取文本，供整批笔记共用。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional, Union

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RESUME_DIR = PROJECT_ROOT / "resume"

DEFAULT_RESUME_CANDIDATES = [
    RESUME_DIR / "resume.docx",
    RESUME_DIR / "简历.docx",
    RESUME_DIR / "resume.pdf",
    RESUME_DIR / "简历.pdf",
]


def find_resume_file(explicit: Optional[Union[Path, str]] = None) -> Path:
    """优先使用显式路径，否则按顺序探测 resume/ 目录下的简历文件。"""
    if explicit:
        path = Path(explicit)
        if not path.exists():
            raise FileNotFoundError(f"指定的简历文件不存在: {path}")
        return path
    for candidate in DEFAULT_RESUME_CANDIDATES:
        if candidate.exists():
            return candidate
    names = "、".join(f"resume/{c.name}" for c in DEFAULT_RESUME_CANDIDATES)
    raise FileNotFoundError(
        f"没有找到简历文件。请把简历放到 resume/ 文件夹，文件名取以下任意一种：{names}；"
        "或使用 --resume 指定简历文件路径"
    )


def read_resume_text(resume_path: Union[Path, str]) -> str:
    """按扩展名读取简历文本；空内容视为扫描版并给出中文提示。"""
    path = Path(resume_path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    raise ValueError(f"不支持的简历格式: {suffix}（仅支持 .docx / .pdf）")


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "未安装 python-docx 库，请先在 Anaconda 虚拟环境中执行 "
            "pip install -r requirements.txt 后重试"
        )
    doc = Document(str(path))
    xml = doc.element.xml
    if "<w:txbxContent>" in xml:
        # 文字放在文本框/形状里：python-docx 的 paragraphs 读不到，
        # 改为按 w:p 段落边界提取全部 w:t 文本；
        # Word 的 mc:AlternateContent 会让同一内容出现两份，按首次出现去重。
        lines: list[str] = []
        seen: set[str] = set()
        for block in re.split(r"</w:p>", xml):
            line = "".join(
                re.findall(r"<w:t[^>]*>([^<]*)</w:t>", block)
            ).strip()
            if line and line not in seen:
                seen.add(line)
                lines.append(line)
        text = "\n".join(lines)
    else:
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    if not text.strip():
        raise RuntimeError("简历 .docx 中没有提取到任何文字，请检查文件内容")
    return text


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError(
            "未安装 pypdf 库，请先在 Anaconda 虚拟环境中执行 "
            "pip install -r requirements.txt 后重试"
        )
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            parts.append(page_text.strip())
    text = "\n".join(parts)
    if not text.strip():
        raise RuntimeError(
            "PDF 中没有提取到任何文字，可能是扫描版/图片版 PDF。"
            "请提供文字版 PDF（可复制的）或改用 .docx 简历"
        )
    return text
