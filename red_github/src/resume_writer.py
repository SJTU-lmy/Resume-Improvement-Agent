# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path
from typing import Dict
from xml.sax.saxutils import escape

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# schema 的 8 个字段
SCHEMA_KEYS = [
    "job_name",
    "base_city",
    "contact_email",
    "job_duty",
    "job_requirement",
    "post_type",
    "resume_optimize_suggestion",
    "target_optimized_resume",
]

_HEADERS = {
    "教育经历", "教育背景", "项目经历", "项目经验", "实习经历", "工作经历",
    "个人优势", "专业技能", "技能证书", "荣誉奖项", "获奖经历", "核心优势",
    "联系方式", "语言能力", "证书", "自我评价",
}


def _import_libs(include_pdf: bool):
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise RuntimeError("未安装 python-docx 库，请先执行 pip install -r requirements.txt")
    if include_pdf:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate
        except ImportError:
            raise RuntimeError("未安装 reportlab 库，请先执行 pip install -r requirements.txt")
        return (Document, Pt, A4, ParagraphStyle, pdfmetrics, UnicodeCIDFont, Paragraph, SimpleDocTemplate)
    return (Document, Pt, None, None, None, None, None, None)


def write_optimized_resume(
    sample_id: str,
    resume_text: str,
    output_dir: Path = PROJECT_ROOT / "new",
    include_pdf: bool = False,
) -> Dict[str, str]:
    """把改写后的简历文本写成 .docx（可选 .pdf），返回已生成文件的路径字典。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    lines = [ln.strip() for ln in str(resume_text).splitlines() if ln.strip()]
    if not lines:
        raise RuntimeError("改写简历文本为空，无法生成文件")

    Document, Pt, A4, ParagraphStyle, pdfmetrics, UnicodeCIDFont, Paragraph, SimpleDocTemplate = _import_libs(include_pdf)

    # --- .docx ---
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    from docx.oxml.ns import qn

    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for line in lines:
        para = doc.add_paragraph()
        is_header = line.rstrip("：:。") in _HEADERS
        segments = line.split("**")
        for idx, seg in enumerate(segments):
            if not seg:
                continue
            run = para.add_run(seg)
            run.bold = is_header or (idx % 2 == 1)
        if is_header and not para.runs:
            para.add_run(line).bold = True
        if is_header:
            para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(0)
    docx_path = output_dir / f"{sample_id}_optimized.docx"
    doc.save(str(docx_path))
    files = {"docx": str(docx_path)}

    if include_pdf:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        cjk_style = ParagraphStyle(
            "CJK", fontName="STSong-Light", fontSize=10, leading=16, wordWrap="CJK"
        )
        pdf_path = output_dir / f"{sample_id}_optimized.pdf"
        pdf_doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=A4,
            topMargin=72, bottomMargin=72, leftMargin=72, rightMargin=72,
        )
        story = [Paragraph(escape(line), cjk_style) for line in lines]
        pdf_doc.build(story)
        files["pdf"] = str(pdf_path)
    return files
